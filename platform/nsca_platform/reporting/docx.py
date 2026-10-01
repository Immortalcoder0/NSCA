"""Editable DOCX export — python-docx, TECH-STACK.md §11, FR-RPT-019,
backend-02 §4.6, M4.

Renders from the **same** `Document` document model as `pdf.py`
(DD-09.1-01: "PDF, DOCX and HTML therefore cannot diverge in content") — no
second pass over `DeviceReportData`, only a second renderer over the
already-composed tree.

Determinism note: python-docx does not stamp `core.xml`'s `created`/
`modified` from the clock unless told to — this module sets them explicitly
from `generated_at` (the same recorded-input principle as pikepdf's
`/CreationDate`, backend-02 §2.5) so a DOCX re-render is byte-comparable
too, even though the milestone's primary V-10.1-001 fixture is the PDF path.
"""

from __future__ import annotations

import io
import zipfile
from datetime import datetime, timezone

from docx import Document as DocxDocument
from docx.shared import Pt

from nsca_platform.reporting.document_model import (
    Chart,
    Document,
    EvidenceExcerpt,
    Heading,
    Manifest,
    Paragraph,
    Playbook,
    SeverityIndicator,
    Table,
)


def _add_block(doc: DocxDocument, block) -> None:
    if isinstance(block, Heading):
        level = min(max(block.level, 1), 4)
        doc.add_heading(block.text, level=level)
    elif isinstance(block, Paragraph):
        p = doc.add_paragraph(block.text)
        if block.kind == "warning":
            for run in p.runs:
                run.bold = True
    elif isinstance(block, Table):
        table = doc.add_table(rows=1, cols=max(len(block.headers), 1))
        table.style = "Light Grid"
        hdr_cells = table.rows[0].cells
        for i, h in enumerate(block.headers):
            hdr_cells[i].text = h
        for row in block.rows:
            cells = table.add_row().cells
            for i, c in enumerate(row.cells):
                if i < len(cells):
                    cells[i].text = c
    elif isinstance(block, EvidenceExcerpt):
        doc.add_paragraph(block.artifact_filename, style="Intense Quote")
        for ln in block.lines:
            marker = ">" if ln.highlighted else " "
            text = f"{marker} {ln.number:>4} | {ln.text}"
            if ln.masked:
                text += " [masked]"
            p = doc.add_paragraph(text)
            p.style.font.name = "Courier New"
            p.style.font.size = Pt(9)
    elif isinstance(block, SeverityIndicator):
        doc.add_paragraph(f"[{block.glyph}] {block.label}")
    elif isinstance(block, Chart):
        doc.add_paragraph(block.title, style="Intense Quote")
        for bar in block.bars:
            filled = 0 if bar.total == 0 else (bar.value * 30) // bar.total
            doc.add_paragraph(f"{bar.label:>10} {'#' * filled}{'.' * (30 - filled)} {bar.value}")
    elif isinstance(block, Playbook):
        for step in block.steps:
            prefix = "[BLOCKED] " if step.blocked else ""
            doc.add_paragraph(f"{prefix}STEP {step.id}  {step.command}  IMPACT: {step.impact.upper()}")
            if step.reason:
                doc.add_paragraph(step.reason)
            if step.rollback:
                incomplete = " (INCOMPLETE)" if step.rollback_incomplete else ""
                doc.add_paragraph(f"Rollback{incomplete}: {step.rollback}")
        for cp in block.checkpoints:
            doc.add_paragraph(f"VERIFY BEFORE CONTINUING (after step {cp.after_step}): {cp.instruction}")
        for control_id, reason in block.unavailable:
            doc.add_paragraph(f"Remediation unavailable — {control_id}: {reason}")
        if block.lockout:
            doc.add_paragraph(f"Lockout analysis: {block.lockout.result} — {block.lockout.reason}")
        if block.dry_run:
            dr = block.dry_run
            doc.add_paragraph(
                f"Dry run: {dr.before_compliance_num}/{dr.before_compliance_den} -> "
                f"{dr.after_compliance_num}/{dr.after_compliance_den}; resolved {len(dr.resolved)}; "
                f"newly failing {len(dr.newly_failing)}"
            )
            doc.add_paragraph(dr.fidelity_note)
    elif isinstance(block, Manifest):
        doc.add_paragraph(f"report_id: {block.report_id}")
        doc.add_paragraph(f"generated: {block.generated_at}")
        for filename, sha in block.inputs:
            doc.add_paragraph(f"  {filename}  sha256:{sha}")
        for entry in block.knowledge_base:
            doc.add_paragraph(f"  {entry.key}: {entry.value}")
        doc.add_paragraph(f"report_hash: {block.report_hash or 'PENDING'}")
        doc.add_paragraph(block.statement)
    else:  # pragma: no cover - exhaustive over document_model.Block
        doc.add_paragraph(str(block))


def _normalise_docx_zip(raw_bytes: bytes, *, generated_at: datetime) -> bytes:
    """DOCX is an OOXML zip container. python-docx's underlying writer
    (Python's `zipfile`) stamps every entry's `date_time` from
    `time.localtime()` at save time — confirmed empirically (this module's
    tests caught it: two identical renders differed at a fixed byte offset
    inside the zip local file header). This is the DOCX-equivalent of the
    `/CreationDate` problem TECH-STACK.md §11 documents for WeasyPrint's PDF
    output, and gets the same treatment: every entry's timestamp is
    rewritten from `generated_at` (never the clock), entries kept in their
    original order and compression settings, so two renders of the same
    `Document` at the same `generated_at` are byte-identical."""

    ts = generated_at if generated_at.tzinfo else generated_at.replace(tzinfo=timezone.utc)
    ts = ts.astimezone(timezone.utc)
    fixed_date_time = (ts.year, ts.month, ts.day, ts.hour, ts.minute, ts.second)

    src = zipfile.ZipFile(io.BytesIO(raw_bytes))
    out_buf = io.BytesIO()
    with zipfile.ZipFile(out_buf, "w", zipfile.ZIP_DEFLATED) as out:
        for info in src.infolist():
            new_info = zipfile.ZipInfo(filename=info.filename, date_time=fixed_date_time)
            new_info.compress_type = info.compress_type
            new_info.external_attr = info.external_attr
            out.writestr(new_info, src.read(info.filename))
    return out_buf.getvalue()


def render_docx(document: Document, *, generated_at: datetime) -> bytes:
    docx = DocxDocument()
    core = docx.core_properties
    ts = generated_at if generated_at.tzinfo else generated_at.replace(tzinfo=timezone.utc)
    core.created = ts
    core.modified = ts
    core.last_modified_by = "NSCA"
    core.author = "NSCA"
    core.revision = 1

    title = f"{document.device_hostname or 'ESTATE'} — Compliance Audit Report"
    docx.add_heading(title, level=0)
    for section in document.sections:
        docx.add_heading(section.title, level=1)
        for block in section.blocks:
            _add_block(docx, block)

    buf = io.BytesIO()
    docx.save(buf)
    return _normalise_docx_zip(buf.getvalue(), generated_at=generated_at)
